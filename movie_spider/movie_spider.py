#!/usr/bin/python3
# -*- coding: utf-8 -*-

"""
This crawler program scan movie org website looking for movie data, write it into a xlsx file.
This crawler program help my wife to make her Academic project application materials.
"""


import DrissionPage
import openpyxl
import openpyxl.styles
import os
import re
import sys


class Spider(object):
    def __init__(self, website='https://www.chinafilm.gov.cn/'):
        self.website = website
        self.page = DrissionPage.SessionPage()

    def get_registration_xlsx_set(self, path='xxgk/gsxx/dybalx/'):
        xlsx_set = set()
        for i in range(25):
            if i == 0:
                self.page.get(url=self.website + path)
            else:
                self.page.get(url=self.website + path + 'index_{}.html'.format(i))
            for element in self.page.eles('@class=m2r_a'):
                xlsx_set.add(element.attr('href'))
        return xlsx_set

    def get_release_xlsx_set(self, path='xxgk/gsxx/dygyxkz/'):
        xlsx_set = set()
        for i in range(4):
            if i == 0:
                self.page.get(url=self.website + path)
            else:
                self.page.get(url=self.website + path + 'index_{}.html'.format(i))
            for element in self.page.eles('@class=m2r_a'):
                xlsx_set.add(element.attr('href'))
        return xlsx_set

    def get_registration_movie_list(self, xlsx_set=set):
        movie_list = list()
        for xlsx in xlsx_set:
            self.page.get(url=xlsx)
            table_list = self.page.eles('@class=hmc4Table')
            for table in table_list:
                if len(table.eles('@tag()=tr')) == 1:
                    continue
                trs = table.eles('@@tag()=tr@!class')
                for tr in trs:
                    _link       = tr.ele('@@target=_blank@@href').attr('href')
                    _no         = tr.ele('@@target=_blank@@href').raw_text
                    try:
                        _year   = re.search(r'.+(\d{4}).+\d+.+', _no).group(1)
                    except AttributeError as err:
                        _year   = None
                    _movie      = tr.ele('@tag()=td').next().next().raw_text.strip()
                    _company_td = tr.ele('@tag()=td').next().next().next().ele('@tag()=script').raw_text
                    _company    = re.search(r'var.+', _company_td).group(0).split('\'')[1]
                    _writer_td  = tr.ele('@tag()=td').next().next().next().next().ele('@tag()=script').raw_text
                    _writer     = re.search(r'var.+', _writer_td).group(0).split('\'')[1]
                    _province   = tr.eles('@tag()=td')[-1].raw_text.strip()
                    movie_list.append({'link'    : _link,
                                       'no'      : _no,
                                       'year'    : _year,
                                       'movie'   : _movie,
                                       'company' : _company,
                                       'writer'  : _writer,
                                       'province': _province})
        return movie_list

    def get_release_movie_list(self, xlsx_set=set):
        movie_list = list()
        for xlsx in xlsx_set:
            self.page.get(url=xlsx)
            table_list = self.page.eles('@class=hmc4Table')
            for table in table_list:
                if len(table.eles('@tag()=tr')) == 1:
                    continue
                trs = table.eles('@@tag()=tr@!class')
                for tr in trs[1:]:    # 需要丢弃数据表的表头
                    _type    = tr.eles('@tag()=td')[1].raw_text.strip()
                    try:
                        _id  = re.search(r'var.+"(\d+)\"', tr.eles('@tag()=td')[2].raw_text.strip()).group(1)
                    except AttributeError as err:
                        _id  = None
                    _no      = tr.eles('@tag()=td')[3].raw_text.strip()
                    _movie   = tr.eles('@tag()=td')[4].raw_text.strip()
                    _company = tr.eles('@tag()=td')[5].raw_text.strip()
                    movie_list.append({'type': _type,
                                       'id': _id,
                                       'no': _no,
                                       'movie': _movie,
                                       'company': _company})
        return movie_list


class Excel(object):
    def __init__(self, filename=str):
        self.filename = filename
        self.excel = self.open_excel()

    def open_excel(self):
        if os.path.isfile(path=self.filename):
            os.remove(self.filename)
        return openpyxl.workbook.Workbook()

    def save_excel(self, encode='UTF-8'):
        if encode == 'UTF-8':
            self.excel.save(filename=self.filename)
        else:
            self.excel.save(filename=self.filename.decode(encode).encode(encode))

    def create_sheet(self, sheetname=str):
        self.excel.create_sheet(title=sheetname, index=0)

    def delete_sheet(self, sheetname=str):
        self.excel.remove(worksheet=self.get_sheet(sheetname=sheetname))

    def get_sheet(self, sheetname=str):
        """
        说明：方法直接返回sheet对象，若sheet不存在则创建sheet对象并返回
        """
        try:
            return self.excel[sheetname]
        except KeyError as err:
            return self.excel.create_sheet(title=sheetname, index=0)

    def set_registration_sheet(self, sheetname=str, movie_list=list):
        sheet = self.excel[sheetname]
        collist = ['备案号', '电影片名', '出品单位', '编剧', '备案省份', '年份', '链接']
        for i in range(len(collist)):
            sheet.cell(row=1, column=(i+1), value=collist[i])
        index = 2
        for movie in movie_list:
            sheet.cell(row=index, column=1, value=movie['no'])
            sheet.cell(row=index, column=2, value=movie['movie'])
            sheet.cell(row=index, column=3, value=movie['company'])
            sheet.cell(row=index, column=4, value=movie['writer'])
            sheet.cell(row=index, column=5, value=movie['province'])
            sheet.cell(row=index, column=6, value=movie['year'])
            sheet.cell(row=index, column=7, value=movie['link'])
            index += 1

    def set_release_sheet(self, sheetname=str, movie_list=list):
        sheet = self.excel[sheetname]
        collist = ['电影类别', '编码', '公映号', '片名', '第一出品单位']
        for i in range(len(collist)):
            sheet.cell(row=1, column=(i + 1), value=collist[i])
        index = 2
        for movie in movie_list:
            sheet.cell(row=index, column=1, value=movie['type'])
            sheet.cell(row=index, column=2, value=movie['no'])
            sheet.cell(row=index, column=3, value=movie['id'])
            sheet.cell(row=index, column=4, value=movie['movie'])
            sheet.cell(row=index, column=5, value=movie['company'])
            index += 1


if __name__ == '__main__':
    spider = Spider()
    registration_xlsx_set = spider.get_registration_xlsx_set()
    registration_movie_list = spider.get_registration_movie_list(xlsx_set=registration_xlsx_set)
    release_xlsx_set = spider.get_release_xlsx_set()
    release_movie_list = spider.get_release_movie_list(xlsx_set=release_xlsx_set)
    xls = Excel(filename='C:/电影目录.xlsx')
    xls.create_sheet(sheetname='电影备案表')
    xls.set_registration_sheet(sheetname='电影备案表', movie_list=registration_movie_list)
    xls.create_sheet(sheetname='电影上映表')
    xls.set_release_sheet(sheetname='电影上映表', movie_list=release_movie_list)
    xls.delete_sheet(sheetname='Sheet')
    xls.save_excel()
